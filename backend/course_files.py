"""Bounded, all-or-nothing employee file imports and scoped, formula-safe exports."""
import csv
import io
import zipfile
import re
from datetime import date, datetime
from uuid import UUID

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from backend.auth import AuthError
from backend.course import EmployeeInput, values, T, SHANGHAI
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from defusedxml.ElementTree import fromstring

COLUMNS = ['employee_no', 'full_name', 'department_id', 'position_id', 'hire_date', 'employment_status']
MAX_BYTES = 2 * 1024 * 1024


def media(kind):
    return 'text/csv; charset=utf-8' if kind == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def safe_cell(value):
    value = str(value) if value is not None else ''
    if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        return "'" + value
    return value


def encode_sheet(rows, kind):
    if kind == 'csv':
        stream = io.StringIO(newline=''); writer = csv.writer(stream)
        writer.writerows([[safe_cell(v) for v in row] for row in rows])
        return stream.getvalue().encode('utf-8-sig')
    book = Workbook(); sheet = book.active; sheet.title = 'employees'
    for row in rows:
        sheet.append([safe_cell(v) for v in row])
    stream = io.BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


def decode_sheet(content, kind):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('file size')
    if kind == 'csv':
        rows = list(csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True))
    else:
        # Refuse zip bombs, macros, external links and formulas before data-only interpretation.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(info.file_size for info in archive.infolist()) > 10 * MAX_BYTES:
                raise ValueError('expanded workbook too large')
            if any('vbaProject' in name or 'externalLinks/' in name for name in archive.namelist()):
                raise ValueError('active workbook content')
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    for element in fromstring(archive.read(name)).iter():
                        if element.tag.endswith('}c') and not re.fullmatch(r'[A-F](?:[1-9]|[1-9]\d|10[01])', element.get('r','')):
                            raise ValueError('actual cell outside template or row limit')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1: raise ValueError('single worksheet required')
            sheet = book.active
            sheet.reset_dimensions()  # Ignore attacker-controlled dimension metadata; inspect actual cells.
            rows = []
            for cells in sheet.iter_rows(max_row=102, max_col=7):
                if any(cell.data_type == 'f' for cell in cells): raise ValueError('formula cell')
                if cells[6].value is not None: raise ValueError('extra column')
                rows.append([cell.value for cell in cells[:6]])
            while rows and all(value is None for value in rows[-1]): rows.pop()
            if len(rows)>101: raise ValueError('row count')
        finally:
            book.close()
    if not rows or rows[0] != COLUMNS or not 1 <= len(rows)-1 <= 100:
        raise ValueError('template columns or row count')
    return rows[1:]


def import_employees(service, session, now, content, kind, mode):
    try:
        rows = decode_sheet(content, kind)
    except (ValueError, UnicodeError, zipfile.BadZipFile, csv.Error, KeyError, TypeError, OSError, ParseError, DefusedXmlException, InvalidFileException):
        return {'valid': False, 'errors': [{'row': 1, 'code': 'INVALID_FILE_TEMPLATE'}], 'imported': 0}
    prepared, errors, seen = [], [], set()
    for number, row in enumerate(rows, 2):
        try:
            if len(row) != len(COLUMNS): raise ValueError('columns')
            data = dict(zip(COLUMNS, row))
            if isinstance(data['hire_date'], datetime): data['hire_date'] = data['hire_date'].date().isoformat()
            payload = EmployeeInput(employee_no=data['employee_no'], full_name=data['full_name'], current_department_id=data['department_id'],
                current_position_id=data['position_id'], hire_date=data['hire_date'], employment_status=data['employment_status'])
            item = values(payload)
            service.active_position(session, item['current_department_id'], item['current_position_id'])
            if item['employment_status'] in ('ACTIVE','PROBATION','EXIT_PROCESSING') and item['hire_date'] > now.astimezone(SHANGHAI).date():
                raise ValueError('future hire')
            if item['employee_no'] in seen or session.execute(select(T['employees'].c.id).where(T['employees'].c.employee_no == item['employee_no'])).first():
                raise ValueError('duplicate')
            seen.add(item['employee_no']); prepared.append(item)
        except (ValueError, TypeError, ValidationError, AuthError):
            errors.append({'row':number, 'code':'INVALID_OR_DUPLICATE_EMPLOYEE'})
    if errors:
        return {'valid':False, 'errors':errors, 'imported':0}
    if mode == 'commit':
        for item in prepared: service.save_employee(session, now, item)
    return {'valid':True, 'errors':[], 'count':len(prepared), 'imported':len(prepared) if mode == 'commit' else 0}


def export_employees(service, session, principal, q, kind):
    table = T['employees']; cond = [service.employee_scope(session, principal)]
    if q:
        from sqlalchemy import or_
        cond.append(or_(table.c.full_name.contains(q, autoescape=True), table.c.employee_no.contains(q, autoescape=True)))
    rows = session.execute(select(table).where(*cond).order_by(table.c.employee_no)).mappings()
    data = [COLUMNS]
    for row in rows:
        data.append([row['employee_no'],row['full_name'],row['current_department_id'],row['current_position_id'],
                     row['hire_date'].isoformat() if row['hire_date'] else '',row['employment_status']])
    return encode_sheet(data, kind)
