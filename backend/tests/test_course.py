import io
import secrets
from datetime import datetime, timezone, timedelta
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from backend.app import create_app
from backend.settings import Settings
from backend.models import metadata
from conftest import Clock


@pytest.fixture
def course(tmp_path):
    from backend.course_demo import create_demo, demo_path
    settings = Settings(demo_path('suite', root=tmp_path), 'https://hrms.test')
    passwords = {role: secrets.token_urlsafe(24) for role in ('admin', 'hr', 'manager', 'employee')}
    clock = Clock()
    clock.value = datetime(2026,10,26,0,59,tzinfo=timezone.utc)  # Monday 08:59 Shanghai
    create_demo(settings, passwords, root=tmp_path, clock=clock)
    app = create_app(settings, clock)
    with TestClient(app, base_url=settings.public_origin, headers={'Origin': settings.public_origin}) as client:
        headers, profiles = {}, {}
        for role, password in passwords.items():
            token = client.post('/api/v1/auth/login', json={'username': 'fixture-' + role, 'password': password}).json()['data']['access_token']
            headers[role] = {'Authorization': 'Bearer ' + token}
            profiles[role] = client.get('/api/v1/auth/me', headers=headers[role]).json()['data']
        yield app, client, clock, headers, profiles


def test_positions_and_employees_real_crud_and_permissions(course):
    app, client, _, headers, profiles = course
    h = headers['hr']
    department = client.get('/api/v1/departments', headers=h).json()['data']['items'][0]['id']
    body = {'department_id': department, 'code': 'P-1', 'name': '课程岗位', 'grade_code': 'G2', 'status': 'ACTIVE'}
    assert client.post('/api/v1/positions', json=body).status_code == 401
    assert client.post('/api/v1/positions', headers=headers['admin'], json=body).status_code == 403
    response = client.post('/api/v1/positions', headers=h, json=body)
    assert response.status_code == 201
    position = response.json()['data']
    assert client.post('/api/v1/positions', headers=h, json=body).status_code == 409
    employee = {'employee_no': 'E-1', 'full_name': '虚构新员工', 'current_department_id': department,
                'current_position_id': position['id'], 'hire_date': '2026-01-01', 'employment_status': 'ACTIVE'}
    response = client.post('/api/v1/employees', headers=h, json=employee)
    assert response.status_code == 201
    eid = response.json()['data']['id']
    assert client.post('/api/v1/employees', headers=h, json=employee).status_code == 409
    assert client.get('/api/v1/employees?q=E-1&page_size=1', headers=h).json()['data']['total'] == 1
    detail = client.get('/api/v1/employees/' + eid, headers=h)
    assert client.patch('/api/v1/employees/' + eid, headers=h, json=employee).status_code == 428
    changed = client.patch('/api/v1/employees/' + eid, headers={**h, 'If-Match': detail.headers['etag']}, json={**employee, 'full_name': '已更名'})
    assert changed.status_code == 200
    assert client.patch('/api/v1/employees/' + eid, headers={**h, 'If-Match': detail.headers['etag']}, json=employee).status_code == 412
    assert client.get('/api/v1/employees/' + eid, headers=headers['employee']).status_code == 403
    pdetail = client.get('/api/v1/positions/' + position['id'], headers=h)
    assert client.patch('/api/v1/positions/' + position['id'], headers={**h, 'If-Match': pdetail.headers['etag']}, json={**body, 'status': 'DISABLED'}).status_code == 409
    archive = client.post('/api/v1/employees/' + eid + '/archive', headers={**h, 'If-Match': changed.headers['etag']}, json={})
    assert archive.status_code == 200
    assert client.get('/api/v1/employees/' + eid, headers=h).json()['data']['employment_status'] == 'TERMINATED'
    with app.state.database.transaction() as session:
        assert session.execute(select(func.count()).select_from(metadata.tables['employee_lifecycle_events'])).scalar_one() >= 2


@pytest.mark.parametrize('kind', ['csv', 'xlsx'])
def test_import_preflight_atomicity_template_and_safe_export(course, kind):
    from openpyxl import Workbook, load_workbook
    app, client, _, headers, _ = course
    h = headers['hr']
    assert client.get('/api/v1/employees/template?format=' + kind, headers=h).status_code == 200
    dep = client.get('/api/v1/departments', headers=h).json()['data']['items'][0]['id']
    pos = client.get('/api/v1/positions', headers=h).json()['data']['items'][0]['id']
    columns = ['employee_no', 'full_name', 'department_id', 'position_id', 'hire_date', 'employment_status']
    rows = [['FILE-1', '=1+1', dep, pos, '2026-01-01', 'ACTIVE'], ['FILE-2', '虚构', dep, 'invalid', '2026-01-01', 'ACTIVE']]
    def content(data):
        if kind == 'csv':
            import csv
            stream = io.StringIO(); writer = csv.writer(stream); writer.writerow(columns); writer.writerows(data)
            return stream.getvalue().encode('utf-8-sig')
        book = Workbook(); sheet = book.active; sheet.append(columns)
        for row in data:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if cell.data_type == 'f': cell.data_type = 's'
        stream = io.BytesIO(); book.save(stream); return stream.getvalue()
    raw = content(rows)
    for mode in ('preflight', 'commit'):
        response = client.post(f'/api/v1/employees/import?format={kind}&mode={mode}', headers={**h, 'Content-Type': 'application/octet-stream'}, content=raw)
        assert response.status_code == 200
        assert response.json()['data']['valid'] is False
        assert response.json()['data']['errors'][0]['row'] == 3
    assert client.get('/api/v1/employees?q=FILE-', headers=h).json()['data']['total'] == 0
    response = client.post(f'/api/v1/employees/import?format={kind}&mode=commit', headers={**h, 'Content-Type': 'application/octet-stream'}, content=content(rows[:1]))
    assert response.json()['data']['imported'] == 1
    export = client.get(f'/api/v1/employees/export?format={kind}&q=FILE-', headers=h)
    assert export.status_code == 200
    if kind == 'csv':
        assert "'=1+1" in export.content.decode('utf-8-sig')
    else:
        sheet = load_workbook(io.BytesIO(export.content), data_only=False).active
        assert sheet.cell(2, 2).data_type != 'f'
    assert client.get(f'/api/v1/employees/export?format={kind}', headers=headers['admin']).status_code == 403


def test_attendance_server_clock_rule_and_own_scope(course):
    _, client, clock, h, profile = course
    assert client.post('/api/v1/attendance/clock-in', json={}).status_code == 401
    assert client.post('/api/v1/attendance/clock-in', headers=h['employee'], json={'employee_id':profile['hr']['employee_id'],'time':'2000-01-01'}).status_code == 422
    assert client.post('/api/v1/attendance/clock-out', headers=h['employee'], json={}).status_code == 409
    clock.advance(minutes=12)  # 09:11, grace boundary is 09:10
    response = client.post('/api/v1/attendance/clock-in', headers=h['employee'], json={})
    assert response.status_code == 200
    assert response.json()['data']['clock_in_at'].startswith('2026-10-26T01:11')
    assert client.post('/api/v1/attendance/clock-in', headers=h['employee'], json={}).status_code == 409
    clock.advance(hours=8)  # 17:11 early out
    assert client.post('/api/v1/attendance/clock-out', headers=h['employee'], json={}).status_code == 200
    data = client.get('/api/v1/attendance?month=2026-10', headers=h['employee']).json()['data']
    assert data['summary']['late_days'] == 1 and data['summary']['early_days'] == 1
    assert data['summary']['complete_days'] == 1
    assert client.get('/api/v1/attendance?month=2026-10&employee_id='+profile['hr']['employee_id'], headers=h['employee']).status_code == 403
    rule={'effective_from':'2026-10-27','weekdays':[0,1,2,3,4],'start_time':'08:30','end_time':'17:30','grace_minutes':5}
    assert client.post('/api/v1/attendance/rules', headers=h['employee'], json=rule).status_code == 403
    assert client.post('/api/v1/attendance/rules', headers=h['hr'], json={**rule,'effective_from':'2026-10-26'}).status_code == 409
    assert client.post('/api/v1/attendance/rules', headers=h['hr'], json=rule).status_code == 201
    assert client.get('/api/v1/attendance/rules',headers=h['hr']).json()['data']['current']['start_time']=='09:00'


def test_leave_decisions_overlap_self_cross_department_and_statistics(course):
    _,client,_,h,profile=course
    body={'start_date':'2026-10-27','end_date':'2026-10-27','leave_type':'ANNUAL','reason':'虚构课程请假'}
    response=client.post('/api/v1/leaves',headers=h['employee'],json=body)
    assert response.status_code==201
    leave=response.json()['data']; path='/api/v1/leaves/'+leave['id']
    assert client.post('/api/v1/leaves',headers=h['employee'],json=body).status_code==409
    assert client.post(path+'/decision',headers={**h['employee'],'If-Match':'"v1"'},json={'status':'APPROVED','opinion':'测试'}).status_code==403
    own=client.post('/api/v1/leaves',headers=h['manager'],json=body).json()['data']
    assert client.post('/api/v1/leaves/'+own['id']+'/decision',headers={**h['manager'],'If-Match':'"v1"'},json={'status':'APPROVED','opinion':'测试'}).status_code==403
    decision=client.post(path+'/decision',headers={**h['manager'],'If-Match':'"v1"'},json={'status':'APPROVED','opinion':'虚构批准'})
    assert decision.status_code==200 and decision.json()['data']['decided_at']
    assert client.post(path+'/decision',headers={**h['manager'],'If-Match':'"v2"'},json={'status':'REJECTED','opinion':'再次审批'}).status_code==409
    assert client.post(path+'/cancel',headers={**h['employee'],'If-Match':'"v2"'},json={}).status_code==409
    summary=client.get('/api/v1/attendance?month=2026-10',headers=h['employee']).json()['data']['summary']
    assert summary['leave_days']==1 and summary['expected_days']==summary['planned_days']-1
    cancel=client.post('/api/v1/leaves/'+own['id']+'/cancel',headers={**h['manager'],'If-Match':'"v1"'},json={})
    assert cancel.status_code==200 and cancel.json()['data']['status']=='CANCELLED'
    # Move the employee into a department the existing manager does not control.
    dep=client.post('/api/v1/departments',headers=h['hr'],json={'code':'OTHER','name':'虚构其他部'}).json()['data']
    pos=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'OTHER','name':'虚构职位','grade_code':'G3'}).json()['data']
    eid=profile['employee']['employee_id']; detail=client.get('/api/v1/employees/'+eid,headers=h['hr'])
    assert client.post('/api/v1/employees/'+eid+'/transfer',headers={**h['hr'],'If-Match':detail.headers['etag']},json={'department_id':dep['id'],'position_id':pos['id']}).status_code==200
    other=client.post('/api/v1/leaves',headers=h['employee'],json={**body,'start_date':'2026-10-28','end_date':'2026-10-28'}).json()['data']
    assert client.post('/api/v1/leaves/'+other['id']+'/decision',headers={**h['manager'],'If-Match':'"v1"'},json={'status':'APPROVED','opinion':'越权'}).status_code==403
    assert client.get('/api/v1/leaves?scope=approval',headers=h['manager']).json()['data']['total']==0


def test_salary_integer_cents_revision_and_scope_dashboard(course):
    app,client,_,h,profile=course
    employee=profile['employee']['employee_id']; other=profile['hr']['employee_id']
    body={'employee_id':employee,'month':'2026-10','amount':'12345.67'}
    assert client.put('/api/v1/salaries',headers=h['admin'],json=body).status_code==403
    assert client.put('/api/v1/salaries',headers=h['employee'],json=body).status_code==403
    for bad in (1.1,'-1','1.234','NaN','1e3'):
        assert client.put('/api/v1/salaries',headers=h['hr'],json={**body,'amount':bad}).status_code==422
    response=client.put('/api/v1/salaries',headers=h['hr'],json=body)
    assert response.status_code==200 and response.json()['data']['amount']=='12345.67'
    sid=response.json()['data']['id']
    assert client.put('/api/v1/salaries',headers=h['hr'],json={**body,'amount':'0'}).status_code==428
    assert client.put('/api/v1/salaries',headers={**h['hr'],'If-Match':'"v1"'},json={**body,'amount':'0'}).json()['data']['amount']=='0.00'
    own=client.get('/api/v1/salaries',headers=h['employee']).json()['data']['items']
    assert len(own)==1 and own[0]['employee_id']==employee
    assert client.get('/api/v1/salaries?employee_id='+other,headers=h['employee']).status_code==403
    assert client.get('/api/v1/salaries',headers=h['admin']).status_code==403
    assert client.get('/api/v1/salaries/'+sid+'/history',headers=h['employee']).status_code==403
    assert [r['amount'] for r in client.get('/api/v1/salaries/'+sid+'/history',headers=h['hr']).json()['data']['items']]==['12345.67','0.00']
    dashboard=client.get('/api/v1/dashboard?month=2026-10',headers=h['hr']).json()['data']
    assert dashboard['employee_count']==3 and len(dashboard['salary_trend'])==6
    assert dashboard['salary_trend'][-1]['amount']=='0.00'
    assert 'salary_trend' not in client.get('/api/v1/dashboard?month=2026-10',headers=h['manager']).json()['data']
    assert client.get('/api/v1/dashboard?month=2026-10',headers=h['admin']).status_code==403
    with app.state.database.transaction() as session:
        table=metadata.tables['course_salary_history']
        assert list(session.execute(select(table.c.amount_cents).order_by(table.c.version_no)).scalars())==[1234567,0]
        audits=session.execute(select(metadata.tables['audit_events'].c.metadata_redacted)).scalars().all()
        assert all('12345.67' not in str(row) for row in audits)


def test_admin_accounts_binding_roles_disable_reset_and_session_revocation(course):
    _,client,_,h,profile=course
    assert client.get('/api/v1/users',headers=h['hr']).status_code==403
    assert client.get('/api/v1/users/employee-options',headers=h['admin']).status_code==200
    listing=client.get('/api/v1/users',headers=h['admin']).json()['data']
    employee=next(r for r in listing['items'] if r['id']==profile['employee']['id'])
    admin=next(r for r in listing['items'] if r['id']==profile['admin']['id'])
    assert 'password_hash' not in str(listing)
    change={'employee_id':employee['employee_id'],'roles':['EMPLOYEE'],'department_ids':[],'status':'DISABLED'}
    assert client.patch('/api/v1/users/'+admin['id'],headers={**h['admin'],'If-Match':'"v1"'},json={**change,'roles':['HR']}).status_code==403
    response=client.patch('/api/v1/users/'+employee['id'],headers={**h['admin'],'If-Match':'"v1"'},json=change)
    assert response.status_code==200
    assert client.get('/api/v1/auth/me',headers=h['employee']).status_code==401
    fresh=secrets.token_urlsafe(24)
    assert client.post('/api/v1/users/'+profile['manager']['id']+'/password',headers={**h['admin'],'If-Match':'"v1"'},json={'password':fresh}).status_code==200
    assert client.get('/api/v1/auth/me',headers=h['manager']).status_code==401
    assert client.post('/api/v1/auth/login',json={'username':'fixture-manager','password':fresh}).status_code==200
    # Creating another independent ADMIN is allowed, but never mutate the acting ADMIN's own role.
    assert client.post('/api/v1/users',headers=h['admin'],json={'username':'new-admin','password':secrets.token_urlsafe(24),'roles':['ADMIN'],'employee_id':None}).status_code==201
    assert client.post('/api/v1/users',headers=h['admin'],json={'username':'unbound-hr','password':secrets.token_urlsafe(24),'roles':['HR'],'employee_id':None}).status_code==409


def test_midnight_shanghai_transfer_keeps_real_hr_employment(course):
    app,client,clock,h,profile=course
    # First transfer a fixture HR after the Shanghai day changes but before UTC midnight.
    clock.value=datetime(2026,10,26,16,1,tzinfo=timezone.utc)  # Oct27 00:01 Shanghai
    dep=client.post('/api/v1/departments',headers=h['hr'],json={'code':'MIDNIGHT','name':'虚构午夜部'}).json()['data']
    pos=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'MIDNIGHT','name':'虚构职位','grade_code':'G1'}).json()['data']
    eid=profile['hr']['employee_id']; detail=client.get('/api/v1/employees/'+eid,headers=h['hr'])
    assert client.post('/api/v1/employees/'+eid+'/transfer',headers={**h['hr'],'If-Match':detail.headers['etag']},json={'department_id':dep['id'],'position_id':pos['id']}).status_code==200
    result=client.post('/api/v1/departments',headers=h['hr'],json={'code':'AFTER-MIDNIGHT','name':'仍有效'})
    assert result.status_code==201


def test_admin_can_disable_archived_employee_account(course):
    _,client,_,h,profile=course
    eid=profile['employee']['employee_id']; detail=client.get('/api/v1/employees/'+eid,headers=h['hr'])
    assert client.post('/api/v1/employees/'+eid+'/archive',headers={**h['hr'],'If-Match':detail.headers['etag']},json={}).status_code==200
    result=client.patch('/api/v1/users/'+profile['employee']['id'],headers={**h['admin'],'If-Match':'"v1"'},json={
        'employee_id':eid,'roles':['EMPLOYEE'],'department_ids':[],'status':'DISABLED'})
    assert result.status_code==200


def test_xlsx_understated_dimensions_validate_all_actual_rows(course):
    import zipfile
    from openpyxl import Workbook
    _,client,_,h,_=course
    dep=client.get('/api/v1/departments',headers=h['hr']).json()['data']['items'][0]['id']
    pos=client.get('/api/v1/positions',headers=h['hr']).json()['data']['items'][0]['id']
    book=Workbook();sheet=book.active
    sheet.append(['employee_no','full_name','department_id','position_id','hire_date','employment_status'])
    # Header claims row2 only; actual invalid row3 MUST reject the entire batch.
    sheet.append(['VISIBLE','虚构',dep,pos,'2026-01-01','ACTIVE'])
    sheet.append(['HIDDEN','虚构','invalid','invalid','2026-01-01','ACTIVE'])
    stream=io.BytesIO();book.save(stream)
    output=io.BytesIO()
    with zipfile.ZipFile(stream) as source,zipfile.ZipFile(output,'w') as target:
        for info in source.infolist():
            content=source.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':content=content.replace(b'ref="A1:F3"',b'ref="A1:F2"')
            target.writestr(info,content)
    result=client.post('/api/v1/employees/import?format=xlsx&mode=commit',headers={**h['hr'],'Content-Type':'application/octet-stream'},content=output.getvalue())
    assert result.status_code==200
    data=result.json()['data']
    assert data['valid'] is False and data['errors'][0]['row']==3
    assert client.get('/api/v1/employees?q=VISIBLE',headers=h['hr']).json()['data']['total']==0
