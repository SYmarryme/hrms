import pytest
from test_course import course


def test_authorized_human_labels_and_utc_timestamps_do_not_broaden_scope(course):
    _,client,_,h,profiles=course
    own=client.get('/api/v1/auth/me',headers=h['employee']).json()['data']
    assert own['employee_name'] and own['employee_no'] and own['employment_valid'] is True
    response=client.post('/api/v1/leaves',headers=h['employee'],json={'start_date':'2026-10-27','end_date':'2026-10-27','leave_type':'ANNUAL','reason':'虚构测试'})
    assert response.status_code==201
    leaves=client.get('/api/v1/leaves?scope=approval',headers=h['manager']).json()['data']['items']
    assert leaves[0]['employee_name'] and leaves[0]['employee_no']
    assert leaves[0]['created_at'].endswith('Z')
    decision=client.post('/api/v1/leaves/'+leaves[0]['id']+'/decision',headers={**h['manager'],'If-Match':'"v1"'},json={'status':'APPROVED','opinion':'虚构批准'})
    assert decision.status_code==200
    own_leaves=client.get('/api/v1/leaves',headers=h['employee']).json()['data']['items']
    assert own_leaves[0]['decided_by_username']=='fixture-manager'
    assert client.get('/api/v1/users/audit-events',headers=h['employee']).status_code==403
    events=client.get('/api/v1/users/audit-events',headers=h['admin']).json()['data']['items']
    assert any(event['actor_username']=='fixture-employee' for event in events)
    assert all(event['occurred_at'].endswith('Z') for event in events)
    salary=client.put('/api/v1/salaries',headers=h['hr'],json={'employee_id':profiles['employee']['employee_id'],'month':'2026-10','amount':'12.34'})
    assert salary.status_code==200
    own_salary=client.get('/api/v1/salaries',headers=h['employee']).json()['data']['items']
    assert own_salary[0]['employee_name']==own['employee_name']
    assert client.get('/api/v1/salaries',headers=h['admin']).status_code==403
    history=client.get('/api/v1/salaries/'+salary.json()['data']['id']+'/history',headers=h['hr']).json()['data']['items']
    assert history[0]['actor_username']=='fixture-hr'


def test_export_uses_all_current_filters_not_just_search_text(course):
    _,client,_,h,_=course
    result=client.get('/api/v1/employees/export?format=csv&status=TERMINATED',headers=h['hr'])
    assert result.status_code==200
    assert len(result.content.decode('utf-8-sig').strip().splitlines())==1
    invalid=client.get('/api/v1/employees/export?format=csv&department_id=invalid',headers=h['hr'])
    assert invalid.status_code==422


def test_attendance_rule_save_and_application_reopen_persists_in_isolated_database(course):
    from pathlib import Path
    from fastapi.testclient import TestClient
    from backend.app import create_app
    from backend.settings import Settings
    app,client,clock,h,_=course
    payload={'effective_from':'2026-10-27','weekdays':[0,1,2,3,4],'start_time':'08:45','end_time':'17:45','grace_minutes':7}
    result=client.post('/api/v1/attendance/rules',headers=h['hr'],json=payload)
    assert result.status_code==201
    path=Path(app.state.database.engine.url.database)
    reopened=create_app(Settings(path,'https://hrms.test'),clock)
    with TestClient(reopened,base_url='https://hrms.test',headers={'Origin':'https://hrms.test'}) as after:
        rows=after.get('/api/v1/attendance/rules',headers=h['hr']).json()['data']['items']
        assert any(row['effective_from']=='2026-10-27' and row['grace_minutes']==7 for row in rows)


def test_background_permission_probe_does_not_keep_idle_sessions_alive(course):
    _,client,clock,h,_=course
    clock.advance(hours=23)
    probe={**h['employee'],'X-HRMS-Session-Probe':'1'}
    assert client.get('/api/v1/auth/me',headers=probe).status_code==200
    clock.advance(hours=1)
    status=client.get('/api/v1/auth/me',headers=probe).status_code
    assert status==401


@pytest.mark.parametrize('kind', ['csv', 'xlsx'])
def test_current_query_export_combines_filters_and_is_not_limited_to_page(course,kind):
    import csv,io
    from openpyxl import load_workbook
    _,client,_,h,profiles=course
    dep=client.get('/api/v1/departments',headers=h['hr']).json()['data']['items'][0]
    pos=client.get('/api/v1/positions',headers=h['hr']).json()['data']['items'][0]
    for index in range(23):
        response=client.post('/api/v1/employees',headers=h['hr'],json={'employee_no':f'EXPORT-{index:02}',
            'full_name':'虚构查询员工','current_department_id':dep['id'],'current_position_id':pos['id'],
            'hire_date':'2026-10-26','employment_status':'ACTIVE' if index<22 else 'PROBATION'})
        assert response.status_code==201
    query=f"q=EXPORT-&department_id={dep['id']}&status=ACTIVE"
    listing=client.get('/api/v1/employees?'+query+'&page_size=20',headers=h['hr']).json()['data']
    assert listing['total']==22 and len(listing['items'])==20
    exported=client.get(f'/api/v1/employees/export?format={kind}&'+query,headers=h['hr'])
    assert exported.status_code==200
    if kind=='csv': rows=list(csv.reader(io.StringIO(exported.content.decode('utf-8-sig'))))
    else: rows=list(load_workbook(io.BytesIO(exported.content),read_only=True).active.values)
    assert len(rows)==23
    assert {row[0] for row in rows[1:]}=={f'EXPORT-{i:02}' for i in range(22)}
    # Export repeats authorization rather than trusting the page's current data.
    own=client.get(f'/api/v1/employees/export?format={kind}',headers=h['employee'])
    assert own.status_code==200
    if kind=='csv': own_rows=list(csv.reader(io.StringIO(own.content.decode('utf-8-sig'))))
    else: own_rows=list(load_workbook(io.BytesIO(own.content),read_only=True).active.values)
    assert len(own_rows)==2 and own_rows[1][0]=='FIXTURE-EMPLOYEE'
    assert client.get(f'/api/v1/employees/export?format={kind}&'+query,headers=h['manager']).status_code==200


def test_server_today_changes_at_shanghai_midnight_without_reusing_yesterday_punch(course):
    from datetime import datetime,timezone
    _,client,clock,h,_=course
    assert client.post('/api/v1/attendance/clock-in',headers=h['employee'],json={}).status_code==200
    before=client.get('/api/v1/attendance?month=2026-10',headers=h['employee']).json()['data']
    assert before['today']=='2026-10-26'
    clock.value=datetime(2026,10,26,16,1,tzinfo=timezone.utc)
    after=client.get('/api/v1/attendance?month=2026-10',headers=h['employee']).json()['data']
    assert after['today']=='2026-10-27'
    assert not any(row['work_date']==after['today'] for row in after['items'])
