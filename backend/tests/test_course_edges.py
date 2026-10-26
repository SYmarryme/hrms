"""Adversarial inputs and durable transaction invariants, all in private temporary fixtures."""
import io
import secrets
from datetime import timedelta
import pytest
from sqlalchemy import select, func, update
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.models import metadata
from test_course import course


@pytest.mark.parametrize('payload',[
    {'code':'','name':'x','grade_code':'G1','department_id':'invalid'},
    {'code':'X','name':'x','grade_code':'G1','department_id':'invalid','status':'ARCHIVED'},
])
def test_position_invalid_inputs(course,payload):
    _,client,_,h,_=course
    assert client.post('/api/v1/positions',headers=h['hr'],json=payload).status_code==422


@pytest.mark.parametrize('kind',['csv','xlsx'])
def test_file_bad_columns_limits_formulas_and_unauthorized(course,kind):
    from openpyxl import Workbook
    _,client,_,h,_=course
    dep=client.get('/api/v1/departments',headers=h['hr']).json()['data']['items'][0]['id']
    pos=client.get('/api/v1/positions',headers=h['hr']).json()['data']['items'][0]['id']
    files=[]
    if kind=='csv':
        files=[b'bad,column\r\na,b\r\n',b'\xff\x00',b'x'* (2*1024*1024+1)]
    else:
        for column in (7,8,26):
            book=Workbook();sheet=book.active
            sheet.append(['employee_no','full_name','department_id','position_id','hire_date','employment_status'])
            sheet.append(['X','虚构',dep,pos,'2026-01-01','ACTIVE']);sheet.cell(2,column,'extra')
            content=io.BytesIO();book.save(content);files.append(content.getvalue())
        book=Workbook();sheet=book.active;sheet.append(['employee_no','full_name','department_id','position_id','hire_date','employment_status']);sheet.append(['X','=1+1','invalid','invalid','2026-01-01','ACTIVE'])
        content=io.BytesIO();book.save(content);files.append(content.getvalue());files.append(b'not-zip')
    for content in files:
        response=client.post(f'/api/v1/employees/import?format={kind}&mode=commit',headers={**h['hr'],'Content-Type':'application/octet-stream'},content=content)
        assert response.status_code in (200,413)
        if response.status_code==200:assert response.json()['data']['valid'] is False
    assert client.post(f'/api/v1/employees/import?format={kind}&mode=commit',headers={**h['employee'],'Content-Type':'application/octet-stream'},content=b'bad').status_code==403


def test_transfer_history_position_consistency_and_terminated_session(course):
    app,client,_,h,profile=course
    eid=profile['employee']['employee_id']
    detail=client.get('/api/v1/employees/'+eid,headers=h['hr'])
    dep=client.post('/api/v1/departments',headers=h['hr'],json={'code':'MOVE','name':'虚构调动'}).json()['data']
    pos=client.get('/api/v1/positions',headers=h['hr']).json()['data']['items'][0]['id']
    assert client.post('/api/v1/employees/'+eid+'/transfer',headers={**h['hr'],'If-Match':detail.headers['etag']},json={'department_id':dep['id'],'position_id':pos}).status_code==409
    assert client.post('/api/v1/employees/'+eid+'/archive',headers={**h['hr'],'If-Match':detail.headers['etag']},json={}).status_code==200
    assert client.get('/api/v1/auth/me',headers=h['employee']).status_code==401
    with app.state.database.transaction() as session:
        assert session.execute(select(metadata.tables['employees'].c.id).where(metadata.tables['employees'].c.id==eid)).first()
        assert session.execute(select(func.count()).select_from(metadata.tables['employee_assignments']).where(metadata.tables['employee_assignments'].c.employee_id==eid)).scalar_one()>=1


def test_missing_punch_and_leave_denominator_zero(course):
    _,client,clock,h,_=course
    # End of the same workday: no submitted punches => one missing day.
    clock.advance(hours=10)
    result=client.get('/api/v1/attendance?month=2026-10',headers=h['employee']).json()['data']
    assert result['summary']['missing_days']==1
    assert client.get('/api/v1/attendance?month=2026-09',headers=h['employee']).json()['data']['summary']['attendance_rate'] is None


def test_sqlite_survives_app_restart_and_audit_failure_rolls_back(course,monkeypatch):
    import backend.course as module
    app,client,clock,h,_=course
    dep=client.get('/api/v1/departments',headers=h['hr']).json()['data']['items'][0]['id']
    payload={'department_id':dep,'code':'DURABLE','name':'持久化虚构职位','grade_code':'G1'}
    response=client.post('/api/v1/positions',headers=h['hr'],json=payload)
    assert response.status_code==201
    assert client.get('/api/v1/users/audit-events',headers=h['employee']).status_code==403
    audit_reply=client.get('/api/v1/users/audit-events',headers=h['admin'])
    assert audit_reply.status_code==200
    assert 'password_hash' not in audit_reply.text and 'metadata_redacted' not in audit_reply.text
    def fail(*args,**kwargs):raise RuntimeError('injected audit failure')
    with monkeypatch.context() as patch:
        patch.setattr(module,'audit',fail)
        with pytest.raises(RuntimeError):client.post('/api/v1/positions',headers=h['hr'],json={**payload,'code':'ROLLBACK'})
    assert client.get('/api/v1/positions?q=ROLLBACK',headers=h['hr']).json()['data']['total']==0
    from backend.settings import Settings
    path=app.state.database.engine.url.database
    app2=create_app(Settings(__import__('pathlib').Path(path),'https://hrms.test'),clock)
    with TestClient(app2,base_url='https://hrms.test',headers={'Origin':'https://hrms.test'}) as second:
        assert second.get('/api/v1/positions?q=DURABLE',headers=h['hr']).json()['data']['total']==1
