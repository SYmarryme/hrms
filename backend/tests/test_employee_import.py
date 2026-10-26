"""Business-facing templates and authoritative all-or-nothing employee imports."""
import csv
import io
from datetime import datetime
import pytest
from openpyxl import load_workbook, Workbook
from sqlalchemy import update

from backend.models import metadata
from test_course import course

HEADERS = ['员工工号', '姓名', '所属部门', '所属职位', '入职日期', '员工状态']


def csv_file(rows, headers=HEADERS):
    stream=io.StringIO(); writer=csv.writer(stream); writer.writerow(headers);writer.writerows(rows)
    return stream.getvalue().encode('utf-8-sig')


def upload(client, headers, content, kind='csv', mode='commit'):
    response=client.post(f'/api/v1/employees/import?format={kind}&mode={mode}',
        headers={**headers,'Content-Type':'application/octet-stream'},content=content)
    assert response.status_code==200
    return response.json()['data']


def catalogue(course):
    _,client,_,h,_=course
    department=client.get('/api/v1/departments',headers=h['hr']).json()['data']['items'][0]
    position=client.get('/api/v1/positions',headers=h['hr']).json()['data']['items'][0]
    return department,position


def options(book,column):
    return [row[0] for row in book['业务选项'].iter_rows(min_row=2,min_col=column,max_col=column,values_only=True) if row[0] is not None]


def test_xlsx_chinese_template_has_database_options_for_every_input_row(course):
    _,client,_,h,_=course
    dep,pos=catalogue(course)
    book=load_workbook(io.BytesIO(client.get('/api/v1/employees/template?format=xlsx',headers=h['hr']).content))
    sheet=book.worksheets[0]
    assert [cell.value for cell in sheet[1]]==HEADERS
    assert dep['name'] in options(book,1) and pos['name'] in options(book,3)
    assert '在职' in options(book,6)
    rules=list(sheet.data_validations.dataValidation)
    for coordinate in ('C2','C51','C101','D2','D51','D101','F2','F101'):
        matching=[rule for rule in rules if coordinate in rule.sqref]
        assert len(matching)==1 and matching[0].type=='list'
        assert matching[0].showDropDown is False and matching[0].showErrorMessage is True
        assert matching[0].formula1.lstrip('=') in book.defined_names
    assert sheet['A2'].number_format=='@' and sheet['E2'].number_format=='yyyy-mm-dd'
    # Fill the last allowed employee row, leaving all preceding rows blank.
    for index,value in enumerate(['TEMPLATE-100','虚构模板员工',dep['name'],pos['name'],datetime(2026,1,1),'在职'],1):
        sheet.cell(101,index,value)
    stream=io.BytesIO();book.save(stream)
    result=upload(client,h['hr'],stream.getvalue(),'xlsx')
    assert result['valid'] is True and result['imported']==1
    row=client.get('/api/v1/employees?q=TEMPLATE-100',headers=h['hr']).json()['data']['items'][0]
    assert row['current_department_id']==dep['id'] and row['current_position_id']==pos['id']


@pytest.mark.parametrize('reference',['name','code'])
def test_chinese_csv_maps_names_or_codes_and_translates_employee_status(course,reference):
    _,client,_,h,_=course
    dep,pos=catalogue(course)
    content=csv_file([['BIZ-1','虚构业务员工',dep[reference],pos[reference],'2026-01-01','试用']])
    checked=upload(client,h['hr'],content,mode='preflight')
    assert checked['valid'] is True and checked['imported']==0
    assert client.get('/api/v1/employees?q=BIZ-1',headers=h['hr']).json()['data']['total']==0
    assert upload(client,h['hr'],content)['imported']==1
    row=client.get('/api/v1/employees?q=BIZ-1',headers=h['hr']).json()['data']['items'][0]
    assert row['employment_status']=='PROBATION' and row['current_department_id']==dep['id']
    template=client.get('/api/v1/employees/template?format=csv',headers=h['hr']).content.decode('utf-8-sig')
    assert next(csv.reader(io.StringIO(template)))==HEADERS


def test_duplicate_names_are_disambiguated_by_business_codes_and_department_scope(course):
    _,client,_,h,_=course
    dep,pos=catalogue(course)
    other=client.post('/api/v1/departments',headers=h['hr'],json={'code':'OTHER','name':dep['name']}).json()['data']
    other_pos=client.post('/api/v1/positions',headers=h['hr'],json={
        'department_id':other['id'],'code':pos['code'],'name':pos['name'],'grade_code':'G1'}).json()['data']
    book=load_workbook(io.BytesIO(client.get('/api/v1/employees/template?format=xlsx',headers=h['hr']).content))
    dep_choices=options(book,1);pos_choices=options(book,3)
    assert len(dep_choices)==len(set(dep_choices))==2
    assert len(pos_choices)==len(set(pos_choices))==2
    chosen_dep=next(value for value in dep_choices if 'OTHER' in value)
    chosen_pos=next(value for value in pos_choices if 'OTHER' in value)
    assert dep['name'] in chosen_dep and pos['code'] in chosen_pos
    data=upload(client,h['hr'],csv_file([['NAMES','虚构重名',chosen_dep,chosen_pos,'2026-01-01','在职']]))
    assert data['valid'] is True
    employee=client.get('/api/v1/employees?q=NAMES',headers=h['hr']).json()['data']['items'][0]
    assert employee['current_department_id']==other['id'] and employee['current_position_id']==other_pos['id']
    ambiguous=upload(client,h['hr'],csv_file([['AMBIGUOUS','虚构重名',dep['name'],pos['name'],'2026-01-01','在职']]))
    assert ambiguous['valid'] is False and ambiguous['errors'][0]['code']=='AMBIGUOUS_DEPARTMENT'
    # A raw position code is unique within the resolved department, even if used elsewhere.
    assert upload(client,h['hr'],csv_file([['SCOPED','虚构范围',other['code'],pos['code'],'2026-01-01','在职']]))['valid'] is True


@pytest.mark.parametrize('problem,expected',[
    ('unknown_department','DEPARTMENT_NOT_FOUND'),('disabled_department','DEPARTMENT_DISABLED'),
    ('disabled_position','POSITION_DISABLED'),('mismatched_position','POSITION_DEPARTMENT_MISMATCH'),
    ('unknown_position','POSITION_NOT_FOUND'),('duplicate','DUPLICATE_EMPLOYEE_NO'),
    ('invalid_date','INVALID_HIRE_DATE'),('invalid_status','INVALID_EMPLOYEE_STATUS'),
])
def test_explicit_row_reasons_and_atomic_batch_rejection(course,problem,expected):
    app,client,_,h,_=course
    dep=client.post('/api/v1/departments',headers=h['hr'],json={'code':'BAD-REF','name':'虚构待校验部'}).json()['data']
    pos=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'BAD-REF','name':'虚构待校验岗','grade_code':'G1'}).json()['data']
    # Valid first row must remain uncommitted when the second row is invalid.
    new_dep=client.post('/api/v1/departments',headers=h['hr'],json={'code':'VALID','name':'虚构有效部'}).json()['data']
    new_pos=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':new_dep['id'],'code':'VALID','name':'虚构有效岗','grade_code':'G1'}).json()['data']
    row=['BAD','虚构错误行',dep['code'],pos['code'],'2026-01-01','在职']
    if problem=='unknown_department':row[2]='不存在的部门'
    if problem=='disabled_department':
        with app.state.database.transaction(write=True) as session:session.execute(update(metadata.tables['departments']).where(metadata.tables['departments'].c.id==dep['id']).values(status='DISABLED'))
    if problem=='disabled_position':
        with app.state.database.transaction(write=True) as session:session.execute(update(metadata.tables['positions']).where(metadata.tables['positions'].c.id==pos['id']).values(status='DISABLED'))
    if problem=='mismatched_position':row[3]=new_pos['name']
    if problem=='unknown_position':row[3]='不存在的职位'
    if problem=='duplicate':row[0]='FIXTURE-EMPLOYEE'
    if problem=='invalid_date':row[4]='2026-02-30'
    if problem=='invalid_status':row[5]='已离职'
    content=csv_file([['GOOD','虚构有效行',new_dep['code'],new_pos['code'],'2026-01-01','在职'],row])
    data=upload(client,h['hr'],content)
    assert data['valid'] is False and data['imported']==0
    issue=next(issue for issue in data['errors'] if issue['row']==3)
    assert issue['code']==expected and issue['message'] and issue['field']
    assert client.get('/api/v1/employees?q=GOOD',headers=h['hr']).json()['data']['total']==0


def test_template_snapshots_do_not_authorize_stale_or_tampered_references(course):
    app,client,_,h,_=course
    dep,pos=catalogue(course)
    extra=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'STALE','name':'虚构过期岗位','grade_code':'G1'}).json()['data']
    book=load_workbook(io.BytesIO(client.get('/api/v1/employees/template?format=xlsx',headers=h['hr']).content))
    for i,value in enumerate(['STALE','虚构陈旧模板',dep['name'],extra['name'],'2026-01-01','在职'],1):book.worksheets[0].cell(2,i,value)
    with app.state.database.transaction(write=True) as session:session.execute(update(metadata.tables['positions']).where(metadata.tables['positions'].c.id==extra['id']).values(status='DISABLED'))
    # Reference sheet values and data validation metadata are editable, never authoritative.
    book['业务选项']['C2']='客户端伪造岗位'
    output=io.BytesIO();book.save(output)
    data=upload(client,h['hr'],output.getvalue(),'xlsx')
    assert data['valid'] is False and data['errors'][0]['code']=='POSITION_DISABLED'


def test_template_excludes_disabled_references_and_contains_no_database_ids(course):
    app,client,_,h,_=course
    dep,pos=catalogue(course)
    off=client.post('/api/v1/departments',headers=h['hr'],json={'code':'OFF','name':'虚构停用部'}).json()['data']
    off_pos=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':off['id'],'code':'OFF','name':'虚构停用部门岗位','grade_code':'G1'}).json()['data']
    disabled=client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'STOPPED','name':'虚构停用岗位','grade_code':'G1','status':'DISABLED'}).json()['data']
    with app.state.database.transaction(write=True) as session:session.execute(update(metadata.tables['departments']).where(metadata.tables['departments'].c.id==off['id']).values(status='DISABLED'))
    book=load_workbook(io.BytesIO(client.get('/api/v1/employees/template?format=xlsx',headers=h['hr']).content))
    assert off['name'] not in options(book,1)
    assert off_pos['name'] not in options(book,3) and disabled['name'] not in options(book,3)
    cells=[cell.value for sheet in book for row in sheet for cell in row]
    assert dep['id'] not in cells and pos['id'] not in cells and off['id'] not in cells


def test_duplicate_within_file_and_legacy_id_validation_remain_atomic(course):
    _,client,_,h,_=course
    dep,pos=catalogue(course)
    row=['REPEAT','虚构重复',dep['code'],pos['code'],'2026-01-01','在职']
    result=upload(client,h['hr'],csv_file([row,row]))
    assert result['errors'][0]['row']==3 and result['errors'][0]['code']=='DUPLICATE_EMPLOYEE_NO'
    assert client.get('/api/v1/employees?q=REPEAT',headers=h['hr']).json()['data']['total']==0
    columns=['employee_no','full_name','department_id','position_id','hire_date','employment_status']
    result=upload(client,h['hr'],csv_file([['STRICT','虚构旧格式',dep['code'],pos['id'],'2026-01-01','ACTIVE']],columns))
    assert result['valid'] is False and result['errors'][0]['code']=='DEPARTMENT_NOT_FOUND'


def test_business_position_names_ambiguous_within_department_are_rejected(course):
    _,client,_,h,_=course
    dep,pos=catalogue(course)
    client.post('/api/v1/positions',headers=h['hr'],json={'department_id':dep['id'],'code':'SECOND','name':pos['name'],'grade_code':'G1'})
    result=upload(client,h['hr'],csv_file([['AMB-POS','虚构',dep['code'],pos['name'],'2026-01-01','在职']]))
    assert result['valid'] is False and result['errors'][0]['code']=='AMBIGUOUS_POSITION'


def test_dropdown_references_support_lists_longer_than_inline_excel_limit(course):
    from sqlalchemy import insert
    from uuid import uuid4
    app,client,clock,h,_=course
    dep,_=catalogue(course)
    with app.state.database.transaction(write=True) as session:
        for index in range(15):
            session.execute(insert(metadata.tables['positions']).values(id=str(uuid4()),department_id=dep['id'],code=f'LONG-{index}',
                name='虚构岗位名称用于覆盖长选项列表'+str(index),grade_code='G1',status='ACTIVE',version_no=1,created_at=clock(),updated_at=clock()))
    book=load_workbook(io.BytesIO(client.get('/api/v1/employees/template?format=xlsx',headers=h['hr']).content))
    assert len(','.join(options(book,3)))>255
    assert len(options(book,3))==16
    position_rule=next(rule for rule in book.worksheets[0].data_validations.dataValidation if 'D101' in rule.sqref)
    destinations=list(book.defined_names[position_rule.formula1.lstrip('=')].destinations)
    assert destinations==[('业务选项','$C$2:$C$17')]
